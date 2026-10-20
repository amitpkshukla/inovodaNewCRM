import io
from typing import List, Dict, Any, Union
import pandas as pd
from PyPDF2 import PdfReader
from pptx import Presentation


def extract_text_from_document(file_bytes: bytes, filename: str) -> str:
    """
    Extracts text content deterministically from PDF or PPTX files.
    """
    text_chunks: List[str] = []
    file_stream = io.BytesIO(file_bytes)
    lower_name = filename.lower()

    if lower_name.endswith(".pdf"):
        reader = PdfReader(file_stream)
        for page_idx, page in enumerate(reader.pages):
            page_text = page.extract_text()
            if page_text:
                text_chunks.append(page_text.strip())

    elif lower_name.endswith((".ppt", ".pptx")):
        prs = Presentation(file_stream)
        for slide_idx, slide in enumerate(prs.slides):
            slide_texts: List[str] = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    slide_texts.append(shape.text.strip())
            if slide_texts:
                text_chunks.append("\n".join(slide_texts))
    else:
        # Fallback to UTF-8 text decoding
        text_chunks.append(file_bytes.decode("utf-8", errors="ignore"))

    return "\n\n".join(text_chunks)


def unpivot_horizontal_remarks(
    excel_source: Union[bytes, str],
    id_vars: List[str] = None,
    remark_columns: List[str] = None,
    output_var_name: str = "Month_Period",
    output_value_name: str = "Remarks",
) -> List[Dict[str, Any]]:
    """
    Parses horizontal Excel workbooks and unpivots them using pd.melt().
    Transforms monthly remark columns (e.g. Jan_Remark, Feb_Remark, Mar_Remark)
    into a normalized vertical time-series schema.
    """
    if isinstance(excel_source, bytes):
        df = pd.read_excel(io.BytesIO(excel_source))
    else:
        df = pd.read_excel(excel_source)

    all_cols = list(df.columns)

    # Detect horizontal remark columns if not explicitly provided
    if not remark_columns:
        remark_columns = [
            c for c in all_cols
            if any(m in c.lower() for m in [
                "jan", "feb", "mar", "apr", "may", "jun",
                "jul", "aug", "sep", "oct", "nov", "dec",
                "remark", "update", "note", "q1", "q2", "q3", "q4"
            ])
        ]

    # Identifier columns (Company, Contact, Designation, etc.)
    if not id_vars:
        id_vars = [c for c in all_cols if c not in remark_columns]

    # Perform unpivoting
    unpivoted_df = pd.melt(
        df,
        id_vars=id_vars,
        value_vars=remark_columns,
        var_name=output_var_name,
        value_name=output_value_name,
    )

    # Drop null or blank remarks
    unpivoted_df = unpivoted_df.dropna(subset=[output_value_name])
    unpivoted_df = unpivoted_df[unpivoted_df[output_value_name].astype(str).str.strip() != ""]

    # Fill NaN values in id columns with empty string
    unpivoted_df = unpivoted_df.fillna("")

    return unpivoted_df.to_dict(orient="records")
